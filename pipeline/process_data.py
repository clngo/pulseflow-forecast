#!/usr/bin/env python3
"""Build the PulseFlow executive forecast from the four supplied exports."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook


REPORTING_DATE = date(2026, 10, 15)
BOARD_TARGET = 12_000_000
INPUT_DIR = Path("data/input")
OUTPUT_DIR = Path("data/output")
INPUT_CONFIG = Path("config/input_files.json")
INPUT_ROLES = (
    "salesforce_current",
    "salesforce_prior",
    "vp_scorecard",
    "carebridge_current",
)

OPEN_SF_STAGES = (
    "Introduction",
    "Technical Validation",
    "Business Approval",
    "Contract Negotiations",
)
SF_STAGE_ORDER = {stage: index for index, stage in enumerate(OPEN_SF_STAGES)}
SF_ALLOWED_STAGES = set(OPEN_SF_STAGES) | {"Closed Lost"}
BILLING_FREQUENCIES = ("Annual", "Monthly")
CAREBRIDGE_STAGES = (
    "Introduction",
    "Qualified to Buy",
    "Decision Maker Bought-In",
    "Contract Sent",
    "Closed Lost",
)
CASE_INSENSITIVE_IDENTITY_FIELDS = {"Health_System", "Company_Name"}
ALLOWED_OVERRIDE_FIELDS = {
    "Fiscal_Period",
    "Health_System",
    "Company_Name",
    "Product",
    "Stage",
    "Deal_Stage",
    "Stage_Probability",
    "Amount",
    "Billing_Frequency",
    "Close_Date",
    "Expected_Close_Date",
    "Sales_Rep",
    "AE_Owner",
    "Days_In_Stage",
}

SF_REQUIRED_COLUMNS = {
    "Opportunity_ID",
    "Fiscal_Period",
    "Health_System",
    "Product",
    "Stage",
    "Stage_Probability",
    "Amount",
    "Billing_Frequency",
    "Close_Date",
    "Sales_Rep",
}
CAREBRIDGE_REQUIRED_COLUMNS = {
    "HubSpot_Deal_ID",
    "Fiscal_Period",
    "Company_Name",
    "Product",
    "Deal_Stage",
    "Stage_Probability",
    "Amount",
    "Billing_Frequency",
    "Expected_Close_Date",
    "AE_Owner",
}


def canonical_salesforce_id(value: str) -> str:
    """Return the case-sensitive 15-character Salesforce matching ID."""
    value = str(value).strip()
    if not re.fullmatch(r"[A-Za-z0-9]{15}(?:[A-Za-z0-9]{3})?", value):
        raise ValueError("Salesforce IDs must contain exactly 15 or 18 letters and digits")
    return value[:15]


def annualize_amount(amount: float, billing_frequency: str) -> float:
    """Convert a monthly or annual source amount to annual ACV."""
    if amount < 0:
        raise ValueError("Amount must be nonnegative")
    if billing_frequency == "Monthly":
        return amount * 12
    if billing_frequency == "Annual":
        return amount
    raise ValueError("Billing frequency must be Monthly or Annual")


def probability_from_gpa(gpa: float) -> float:
    """Apply the exact Rule 2 GPA probability boundaries."""
    if gpa >= 3.20:
        return 0.75
    if gpa >= 2.40:
        return 0.50
    if gpa >= 1.60:
        return 0.30
    return 0.10


def calculate_top8_probability(leadership: int, commercial: int, contractual: int) -> tuple[float, float]:
    for score in (leadership, commercial, contractual):
        if not isinstance(score, int) or isinstance(score, bool) or not 1 <= score <= 4:
            raise ValueError("Top 8 scores must be whole numbers from 1 through 4")
    gpa = round((0.20 * leadership) + (0.20 * commercial) + (0.60 * contractual), 2)
    return gpa, probability_from_gpa(gpa)


def normalize_name(value: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", str(value).lower()))


def canonicalize_case(value: str, allowed_values: tuple[str, ...] | set[str]) -> str:
    """Return the official spelling when a known value differs only by case."""
    match = next((allowed for allowed in allowed_values if allowed.casefold() == value.casefold()), None)
    return match if match is not None else value


def comparison_value(record: dict[str, Any], field: str) -> Any:
    value = record.get(field)
    if field in CASE_INSENSITIVE_IDENTITY_FIELDS and isinstance(value, str):
        return value.casefold()
    return value


def name_similarity(left: str, right: str) -> float:
    return round(SequenceMatcher(None, normalize_name(left), normalize_name(right)).ratio(), 3)


def as_number(value: Any, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric") from exc
    if pd.isna(number):
        raise ValueError(f"{label} must be numeric")
    return number


def read_csv(path: Path, required_columns: set[str]) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    missing = sorted(required_columns - set(frame.columns))
    if missing:
        raise ValueError(f"{path.name} is missing required columns: {', '.join(missing)}")
    frame = frame.copy()
    frame["source_row"] = range(2, len(frame) + 2)
    return frame


def read_scorecard(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [sheet.cell(row=4, column=column).value for column in range(1, 12)]
    rows: list[dict[str, Any]] = []
    for row_number in range(5, sheet.max_row + 1):
        values = [sheet.cell(row=row_number, column=column).value for column in range(1, 12)]
        if not any(value not in (None, "") for value in values):
            continue
        record = dict(zip(headers, values))
        record["source_row"] = row_number
        rows.append(record)
    workbook.close()
    return rows


def load_resolutions(root: Path) -> dict[str, Any]:
    path = root / "config" / "resolutions.json"
    if not path.exists():
        return {"scorecard_mappings": [], "duplicate_selections": [], "record_overrides": []}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("config/resolutions.json must contain a JSON object")
    data.setdefault("scorecard_mappings", [])
    data.setdefault("duplicate_selections", [])
    data.setdefault("record_overrides", [])
    if not isinstance(data["scorecard_mappings"], list):
        raise ValueError("scorecard_mappings must be a JSON array")
    if not isinstance(data["duplicate_selections"], list):
        raise ValueError("duplicate_selections must be a JSON array")
    if not isinstance(data["record_overrides"], list):
        raise ValueError("record_overrides must be a JSON array")
    return data


def load_input_files(root: Path) -> dict[str, str]:
    path = root / INPUT_CONFIG
    if not path.is_file():
        raise ValueError(f"Required input configuration is missing: {INPUT_CONFIG}")
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"{INPUT_CONFIG} must contain a JSON object")
    selected: dict[str, str] = {}
    for role in INPUT_ROLES:
        filename = data.get(role)
        if not isinstance(filename, str) or not filename.strip():
            raise ValueError(f"{INPUT_CONFIG} must define a filename for {role}")
        filename = filename.strip()
        if Path(filename).name != filename:
            raise ValueError(f"{role} must be a filename inside {INPUT_DIR}")
        selected[role] = filename
    return selected


def apply_record_overrides(
    frame: pd.DataFrame,
    source_file: str,
    id_column: str,
    overrides: list[dict[str, Any]],
    add_audit: Callable[..., None],
    add_issue: Callable[..., None],
    used_overrides: set[int],
) -> pd.DataFrame:
    frame = frame.copy()
    for position, override in enumerate(overrides):
        if not isinstance(override, dict) or override.get("source") != source_file:
            continue
        used_overrides.add(position)
        record_id = str(override.get("record_id", "")).strip()
        field = str(override.get("field", "")).strip()
        reason = str(override.get("reason", "")).strip()
        source_row = override.get("source_row")
        if field not in frame.columns or field not in ALLOWED_OVERRIDE_FIELDS:
            add_issue(
                "invalid_record_override",
                source_file,
                [],
                record_id,
                f"Override field {field!r} is not permitted",
                override,
                "Choose a supported source field",
            )
            continue
        mask = frame[id_column].astype(str).str.strip().eq(record_id)
        if source_row is not None:
            mask &= frame["source_row"].eq(source_row)
        matches = frame.index[mask].tolist()
        if len(matches) != 1 or not reason:
            add_issue(
                "invalid_record_override",
                source_file,
                [source_row] if source_row is not None else [],
                record_id,
                "A record override must match exactly one row and include a reason",
                {"match_count": len(matches), "override": override},
                "Correct the override target or add its reason",
            )
            continue
        index = matches[0]
        original = frame.at[index, field]
        frame.at[index, field] = override.get("value")
        add_audit(
            "manual_record_override",
            "human_resolved",
            source_file,
            [int(frame.at[index, "source_row"])],
            record_id,
            {"field": field, "value": original},
            {"field": field, "value": override.get("value")},
            f"Applied manual override: {reason}",
            "config/resolutions.json",
            None,
        )
    return frame


def normalize_salesforce(
    frame: pd.DataFrame,
    source_file: str,
    add_audit: Callable[..., None],
    add_issue: Callable[..., None],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    string_fields = [
        "Opportunity_ID",
        "Fiscal_Period",
        "Health_System",
        "Product",
        "Stage",
        "Billing_Frequency",
        "Close_Date",
        "Sales_Rep",
        "Region",
    ]
    for raw in frame.to_dict(orient="records"):
        row = dict(raw)
        for field in string_fields:
            if field in row:
                row[field] = str(row[field]).strip()
        source_row = int(row["source_row"])
        source_id = row["Opportunity_ID"]
        for field, allowed_values in (
            ("Stage", SF_ALLOWED_STAGES),
            ("Billing_Frequency", {"Annual"}),
        ):
            original = row[field]
            row[field] = canonicalize_case(original, allowed_values)
            if row[field] != original:
                add_audit(
                    "capitalization_normalization",
                    "automatically_resolved",
                    source_file,
                    [source_row],
                    source_id,
                    {"field": field, "value": original},
                    {"field": field, "value": row[field]},
                    f"Standardized {field} capitalization",
                    "case_insensitive_category_rule",
                    0,
                )
        errors: list[str] = []
        try:
            canonical_id = canonical_salesforce_id(source_id)
        except ValueError as exc:
            canonical_id = ""
            errors.append(str(exc))
        try:
            amount = as_number(row["Amount"], "Amount")
            if amount < 0:
                raise ValueError("Amount must be nonnegative")
        except ValueError as exc:
            amount = 0.0
            errors.append(str(exc))
        try:
            probability = as_number(row["Stage_Probability"], "Stage_Probability")
            if not 0 <= probability <= 1:
                raise ValueError("Stage_Probability must be between 0 and 1")
        except ValueError as exc:
            probability = 0.0
            errors.append(str(exc))
        if row["Billing_Frequency"] != "Annual":
            errors.append("Salesforce Billing_Frequency must be Annual")
        if row["Stage"] not in SF_ALLOWED_STAGES:
            errors.append(f"Unknown Salesforce stage: {row['Stage']!r}")
        if not row["Fiscal_Period"]:
            errors.append("Fiscal_Period is required")
        if errors:
            add_issue(
                "invalid_salesforce_record",
                source_file,
                [source_row],
                source_id,
                "; ".join(errors),
                row,
                "Correct the source record or add a valid manual override",
            )
            continue
        row.update(
            {
                "source_file": source_file,
                "source_id": source_id,
                "canonical_id": canonical_id,
                "amount": amount,
                "annual_acv": amount,
                "stage_probability": probability,
            }
        )
        records.append(row)
    return records


def normalize_carebridge(
    frame: pd.DataFrame,
    source_file: str,
    add_audit: Callable[..., None],
    add_issue: Callable[..., None],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    string_fields = [
        "HubSpot_Deal_ID",
        "Fiscal_Period",
        "Company_Name",
        "Product",
        "Deal_Stage",
        "Billing_Frequency",
        "Expected_Close_Date",
        "AE_Owner",
    ]
    for raw in frame.to_dict(orient="records"):
        row = dict(raw)
        for field in string_fields:
            row[field] = str(row[field]).strip()
        source_row = int(row["source_row"])
        source_id = row["HubSpot_Deal_ID"]
        for field, allowed_values in (
            ("Deal_Stage", CAREBRIDGE_STAGES),
            ("Billing_Frequency", BILLING_FREQUENCIES),
        ):
            original = row[field]
            row[field] = canonicalize_case(original, allowed_values)
            if row[field] != original:
                add_audit(
                    "capitalization_normalization",
                    "automatically_resolved",
                    source_file,
                    [source_row],
                    source_id,
                    {"field": field, "value": original},
                    {"field": field, "value": row[field]},
                    f"Standardized {field} capitalization",
                    "case_insensitive_category_rule",
                    0,
                )
        errors: list[str] = []
        if not source_id:
            errors.append("HubSpot_Deal_ID is required")
        try:
            amount = as_number(row["Amount"], "Amount")
            annual_acv = annualize_amount(amount, row["Billing_Frequency"])
        except ValueError as exc:
            amount = 0.0
            annual_acv = 0.0
            errors.append(str(exc))
        try:
            probability = as_number(row["Stage_Probability"], "Stage_Probability")
            if not 0 <= probability <= 1:
                raise ValueError("Stage_Probability must be between 0 and 1")
        except ValueError as exc:
            probability = 0.0
            errors.append(str(exc))
        if not row["Fiscal_Period"]:
            errors.append("Fiscal_Period is required")
        if not row["Deal_Stage"]:
            errors.append("Deal_Stage is required")
        if errors:
            add_issue(
                "invalid_carebridge_record",
                source_file,
                [source_row],
                source_id,
                "; ".join(errors),
                row,
                "Correct the source record or add a valid manual override",
            )
            continue
        row.update(
            {
                "source_file": source_file,
                "source_id": source_id,
                "canonical_id": source_id,
                "amount": amount,
                "annual_acv": annual_acv,
                "stage_probability": probability,
            }
        )
        records.append(row)
    return records


def deduplicate_records(
    records: list[dict[str, Any]],
    source_file: str,
    comparison_fields: list[str],
    duplicate_resolutions: list[dict[str, Any]],
    add_audit: Callable[..., None],
    add_issue: Callable[..., None],
    used_duplicate_resolutions: set[int],
) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        groups[record["canonical_id"]].append(record)

    output: list[dict[str, Any]] = []
    for canonical_id, candidates in groups.items():
        if len(candidates) == 1:
            output.append(candidates[0])
            continue

        unique: list[dict[str, Any]] = []
        fingerprints: set[tuple[Any, ...]] = set()
        for candidate in candidates:
            fingerprint = tuple(comparison_value(candidate, field) for field in comparison_fields)
            if fingerprint not in fingerprints:
                fingerprints.add(fingerprint)
                unique.append(candidate)

        if len(unique) == 1:
            retained = unique[0]
            add_audit(
                "exact_duplicate",
                "automatically_resolved",
                source_file,
                [int(candidate["source_row"]) for candidate in candidates],
                retained["source_id"],
                {"duplicate_count": len(candidates)},
                {"calculation_record_count": 1, "retained_source_row": int(retained["source_row"])},
                f"Removed {len(candidates) - 1} identical duplicate row(s)",
                "pipeline_rule",
                0,
            )
            output.append(retained)
            continue

        matching_resolution: tuple[int, dict[str, Any]] | None = None
        source_ids = {candidate["source_id"] for candidate in unique} | {canonical_id}
        for position, resolution in enumerate(duplicate_resolutions):
            if position in used_duplicate_resolutions or not isinstance(resolution, dict):
                continue
            if (
                str(resolution.get("source", "")).strip() == source_file
                and str(resolution.get("opportunity_id", "")).strip() in source_ids
            ):
                matching_resolution = (position, resolution)
                break

        if matching_resolution:
            position, resolution = matching_resolution
            used_duplicate_resolutions.add(position)
            selected_row = resolution.get("selected_source_row")
            reason = str(resolution.get("reason", "")).strip()
            selected = [candidate for candidate in unique if candidate["source_row"] == selected_row]
            if len(selected) != 1 or not reason:
                add_issue(
                    "invalid_duplicate_resolution",
                    source_file,
                    [int(candidate["source_row"]) for candidate in unique],
                    candidates[0]["source_id"],
                    "Duplicate selection must identify one current source row and include a reason",
                    resolution,
                    "Update the duplicate selection in config/resolutions.json",
                )
                continue
            retained = dict(selected[0])
            retained["_resolved_variants"] = unique
            add_audit(
                "conflicting_duplicate",
                "human_resolved",
                source_file,
                [int(candidate["source_row"]) for candidate in unique],
                retained["source_id"],
                {"variants": [compact_variant(candidate) for candidate in unique]},
                {"selected_source_row": int(retained["source_row"])},
                f"Used the authoritatively selected source row: {reason}",
                "config/resolutions.json",
                None,
            )
            output.append(retained)
            continue

        add_issue(
            "conflicting_duplicate",
            source_file,
            [int(candidate["source_row"]) for candidate in unique],
            candidates[0]["source_id"],
            "Duplicate rows contain different meaningful values",
            [{field: candidate.get(field) for field in comparison_fields} for candidate in unique],
            "Select one authoritative source row in config/resolutions.json",
        )
    return output


def compact_variant(record: dict[str, Any]) -> dict[str, Any]:
    stage = record.get("Stage", record.get("Deal_Stage"))
    return {
        "source_row": int(record["source_row"]),
        "stage": stage,
        "stage_probability": record["stage_probability"],
        "amount": record["amount"],
        "annual_acv": record["annual_acv"],
        "expected_value": record["annual_acv"] * record["stage_probability"],
    }


def is_active_salesforce(record: dict[str, Any]) -> bool:
    return record["Fiscal_Period"] == "Q4-2026" and record["Stage"] in OPEN_SF_STAGES


def is_active_carebridge(record: dict[str, Any]) -> bool:
    return record["Fiscal_Period"] == "Q4-2026" and record["Deal_Stage"] != "Closed Lost"


def build_week_over_week(
    current: list[dict[str, Any]], prior: list[dict[str, Any]]
) -> dict[str, list[dict[str, Any]]]:
    current_by_id = {record["canonical_id"]: record for record in current}
    prior_by_id = {record["canonical_id"]: record for record in prior}

    new_deals: list[dict[str, Any]] = []
    slipped: list[dict[str, Any]] = []
    closed_lost: list[dict[str, Any]] = []
    stage_moves: list[dict[str, Any]] = []
    acv_changes: list[dict[str, Any]] = []

    for canonical_id, record in current_by_id.items():
        prior_record = prior_by_id.get(canonical_id)
        base = {
            "opportunity_id": record["source_id"],
            "health_system": record["Health_System"],
            "product": record["Product"],
            "sales_rep": record["Sales_Rep"],
        }
        if not prior_record:
            new_deals.append(
                {
                    **base,
                    "fiscal_period": record["Fiscal_Period"],
                    "stage": record["Stage"],
                    "annual_acv": record["annual_acv"],
                }
            )
            continue

        if prior_record["Fiscal_Period"] == "Q4-2026" and record["Fiscal_Period"] == "Q1-2027":
            slipped.append(
                {
                    **base,
                    "prior_fiscal_period": prior_record["Fiscal_Period"],
                    "current_fiscal_period": record["Fiscal_Period"],
                    "annual_acv": record["annual_acv"],
                }
            )
        if prior_record["Stage"] != "Closed Lost" and record["Stage"] == "Closed Lost":
            closed_lost.append(
                {
                    **base,
                    "prior_stage": prior_record["Stage"],
                    "current_stage": record["Stage"],
                    "annual_acv": record["annual_acv"],
                }
            )
        if is_active_salesforce(record):
            prior_stage = prior_record["Stage"]
            current_stage = record["Stage"]
            if prior_stage in SF_STAGE_ORDER and current_stage in SF_STAGE_ORDER and prior_stage != current_stage:
                direction = "advanced" if SF_STAGE_ORDER[current_stage] > SF_STAGE_ORDER[prior_stage] else "regressed"
                stage_moves.append(
                    {
                        **base,
                        "prior_stage": prior_stage,
                        "current_stage": current_stage,
                        "direction": direction,
                        "prior_probability": prior_record["stage_probability"],
                        "current_probability": record["stage_probability"],
                        "expected_value_change": record["annual_acv"]
                        * (record["stage_probability"] - prior_record["stage_probability"]),
                    }
                )
            if record["annual_acv"] != prior_record["annual_acv"]:
                difference = record["annual_acv"] - prior_record["annual_acv"]
                acv_changes.append(
                    {
                        **base,
                        "prior_annual_acv": prior_record["annual_acv"],
                        "current_annual_acv": record["annual_acv"],
                        "change": difference,
                        "direction": "increase" if difference > 0 else "decrease",
                    }
                )

    sort_key = lambda item: (item["health_system"], item["opportunity_id"])
    return {
        "new_deals": sorted(new_deals, key=sort_key),
        "slipped_deals": sorted(slipped, key=sort_key),
        "closed_lost_deals": sorted(closed_lost, key=sort_key),
        "stage_moves": sorted(stage_moves, key=sort_key),
        "acv_changes": sorted(acv_changes, key=sort_key),
    }


def build_outputs(
    root: Path,
    resolutions_override: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], dict[str, Any] | None]:
    input_files = load_input_files(root)
    current_file = input_files["salesforce_current"]
    prior_file = input_files["salesforce_prior"]
    scorecard_file = input_files["vp_scorecard"]
    carebridge_file = input_files["carebridge_current"]

    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    audits: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    audit_counter = 0
    issue_counter = 0

    def add_audit(
        issue_type: str,
        status: str,
        source: str,
        source_rows: list[int],
        record_id: str | None,
        original_value: Any,
        resulting_value: Any,
        action: str,
        resolution_source: str,
        forecast_impact: Any,
    ) -> None:
        nonlocal audit_counter
        audit_counter += 1
        audits.append(
            {
                "audit_id": f"audit-{audit_counter:03d}",
                "type": issue_type,
                "status": status,
                "source": source,
                "source_rows": source_rows,
                "record_id": record_id,
                "original_value": original_value,
                "resulting_value": resulting_value,
                "action": action,
                "resolution_source": resolution_source,
                "forecast_impact": forecast_impact,
            }
        )

    def add_issue(
        issue_type: str,
        source: str,
        source_rows: list[int],
        record_id: str | None,
        message: str,
        evidence: Any,
        required_action: str,
    ) -> None:
        nonlocal issue_counter
        issue_counter += 1
        issues.append(
            {
                "issue_id": f"review-{issue_counter:03d}",
                "type": issue_type,
                "status": "blocking",
                "source": source,
                "source_rows": source_rows,
                "record_id": record_id,
                "message": message,
                "evidence": evidence,
                "required_action": required_action,
            }
        )

    def build_review() -> dict[str, Any]:
        return {
            "schema_version": 1,
            "reporting_date": REPORTING_DATE.isoformat(),
            "generated_at": generated_at,
            "status": "blocked" if issues else "clear",
            "blocking_issue_count": len(issues),
            "issues": issues,
        }

    source_specs = (
        (current_file, SF_REQUIRED_COLUMNS),
        (prior_file, SF_REQUIRED_COLUMNS),
        (carebridge_file, CAREBRIDGE_REQUIRED_COLUMNS),
        (scorecard_file, None),
    )
    for filename, _ in source_specs:
        if not (root / INPUT_DIR / filename).is_file():
            add_issue(
                "missing_source_file",
                filename,
                [],
                None,
                f"Required source file is missing: {filename}",
                {"missing_file": filename},
                "Restore the required source file and rerun the pipeline",
            )
    if issues:
        return build_review(), None

    for filename, required_columns in source_specs:
        if required_columns is None:
            continue
        path = root / INPUT_DIR / filename
        columns = set(pd.read_csv(path, dtype=str, keep_default_na=False, nrows=0).columns)
        missing = sorted(required_columns - columns)
        if missing:
            add_issue(
                "missing_required_columns",
                filename,
                [],
                None,
                f"{filename} is missing required columns: {', '.join(missing)}",
                {"missing_columns": missing},
                "Restore the required columns and rerun the pipeline",
            )
    if issues:
        return build_review(), None

    resolutions = resolutions_override if resolutions_override is not None else load_resolutions(root)
    resolutions.setdefault("scorecard_mappings", [])
    resolutions.setdefault("duplicate_selections", [])
    resolutions.setdefault("record_overrides", [])

    current_frame = read_csv(root / INPUT_DIR / current_file, SF_REQUIRED_COLUMNS)
    prior_frame = read_csv(root / INPUT_DIR / prior_file, SF_REQUIRED_COLUMNS)
    carebridge_frame = read_csv(root / INPUT_DIR / carebridge_file, CAREBRIDGE_REQUIRED_COLUMNS)

    used_overrides: set[int] = set()
    current_frame = apply_record_overrides(
        current_frame,
        current_file,
        "Opportunity_ID",
        resolutions["record_overrides"],
        add_audit,
        add_issue,
        used_overrides,
    )
    prior_frame = apply_record_overrides(
        prior_frame,
        prior_file,
        "Opportunity_ID",
        resolutions["record_overrides"],
        add_audit,
        add_issue,
        used_overrides,
    )
    carebridge_frame = apply_record_overrides(
        carebridge_frame,
        carebridge_file,
        "HubSpot_Deal_ID",
        resolutions["record_overrides"],
        add_audit,
        add_issue,
        used_overrides,
    )
    for position, override in enumerate(resolutions["record_overrides"]):
        if position not in used_overrides:
            add_issue(
                "stale_record_override",
                str(override.get("source", "config/resolutions.json")),
                [],
                str(override.get("record_id", "")),
                "Record override does not target a supported current source file",
                override,
                "Remove or update the stale override",
            )

    current_records = normalize_salesforce(current_frame, current_file, add_audit, add_issue)
    prior_records = normalize_salesforce(prior_frame, prior_file, add_audit, add_issue)
    carebridge_records = normalize_carebridge(
        carebridge_frame,
        carebridge_file,
        add_audit,
        add_issue,
    )

    sf_comparison = [
        "Fiscal_Period",
        "Health_System",
        "Product",
        "Stage",
        "stage_probability",
        "amount",
        "Billing_Frequency",
        "Close_Date",
        "Created_Date",
        "Days_In_Stage",
        "Sales_Rep",
        "Region",
    ]
    hs_comparison = [
        "Fiscal_Period",
        "Company_Name",
        "Product",
        "Deal_Stage",
        "stage_probability",
        "amount",
        "Billing_Frequency",
        "Expected_Close_Date",
        "AE_Owner",
    ]
    used_duplicate_resolutions: set[int] = set()
    for position, resolution in enumerate(resolutions["duplicate_selections"]):
        if not isinstance(resolution, dict):
            used_duplicate_resolutions.add(position)
            add_issue(
                "invalid_duplicate_resolution",
                "config/resolutions.json",
                [],
                None,
                "Each duplicate selection must be a JSON object",
                resolution,
                "Replace the invalid duplicate selection",
            )
            continue
        selected_row = resolution.get("selected_source_row")
        if (
            not str(resolution.get("source", "")).strip()
            or not str(resolution.get("opportunity_id", "")).strip()
            or not isinstance(selected_row, int)
            or isinstance(selected_row, bool)
            or not str(resolution.get("reason", "")).strip()
        ):
            used_duplicate_resolutions.add(position)
            add_issue(
                "invalid_duplicate_resolution",
                "config/resolutions.json",
                [selected_row] if isinstance(selected_row, int) and not isinstance(selected_row, bool) else [],
                str(resolution.get("opportunity_id", "")).strip() or None,
                "A duplicate selection requires source, opportunity_id, selected_source_row, and reason",
                resolution,
                "Correct the duplicate selection in config/resolutions.json",
            )
    current_records = deduplicate_records(
        current_records,
        current_file,
        sf_comparison,
        resolutions["duplicate_selections"],
        add_audit,
        add_issue,
        used_duplicate_resolutions,
    )
    prior_records = deduplicate_records(
        prior_records,
        prior_file,
        sf_comparison,
        resolutions["duplicate_selections"],
        add_audit,
        add_issue,
        used_duplicate_resolutions,
    )
    carebridge_records = deduplicate_records(
        carebridge_records,
        carebridge_file,
        hs_comparison,
        resolutions["duplicate_selections"],
        add_audit,
        add_issue,
        used_duplicate_resolutions,
    )
    for position, resolution in enumerate(resolutions["duplicate_selections"]):
        if position not in used_duplicate_resolutions:
            add_issue(
                "stale_duplicate_resolution",
                "config/resolutions.json",
                [],
                str(resolution.get("opportunity_id", "")),
                "Duplicate selection no longer matches a current conflicting duplicate",
                resolution,
                "Remove or reconfirm the stale duplicate selection",
            )

    scorecard = read_scorecard(root / INPUT_DIR / scorecard_file)
    scorecard_mappings = resolutions["scorecard_mappings"]
    scorecard_mappings_by_row: dict[int, tuple[int, dict[str, Any]]] = {}
    valid_scorecard_mapping_positions: set[int] = set()
    used_scorecard_mapping_positions: set[int] = set()
    for position, mapping in enumerate(scorecard_mappings):
        if not isinstance(mapping, dict):
            add_issue(
                "invalid_resolution",
                "config/resolutions.json",
                [],
                None,
                "Each scorecard mapping must be a JSON object",
                mapping,
                "Replace the invalid scorecard mapping",
            )
            continue
        source_row = mapping.get("source_row")
        opportunity_id = str(mapping.get("opportunity_id", "")).strip()
        reason = str(mapping.get("reason", "")).strip()
        if (
            not isinstance(source_row, int)
            or isinstance(source_row, bool)
            or source_row < 5
            or not opportunity_id
            or not reason
        ):
            add_issue(
                "invalid_resolution",
                "config/resolutions.json",
                [source_row] if isinstance(source_row, int) and not isinstance(source_row, bool) else [],
                opportunity_id or None,
                "A scorecard mapping requires source_row, opportunity_id, and reason",
                mapping,
                "Correct the scorecard mapping in config/resolutions.json",
            )
            continue
        if source_row in scorecard_mappings_by_row:
            add_issue(
                "invalid_resolution",
                "config/resolutions.json",
                [source_row],
                opportunity_id,
                "Only one scorecard mapping is allowed for each source row",
                mapping,
                "Remove the duplicate scorecard mapping",
            )
            continue
        scorecard_mappings_by_row[source_row] = (position, mapping)
        valid_scorecard_mapping_positions.add(position)

    top8_by_canonical: dict[str, dict[str, Any]] = {}

    current_by_canonical = {record["canonical_id"]: record for record in current_records}
    current_by_source_id = {record["source_id"]: record for record in current_records}

    for score_row in scorecard:
        shorthand = str(score_row.get("Hospital_Shorthand_Name") or "").strip()
        product = str(score_row.get("Product") or "").strip()
        sales_rep = str(score_row.get("Sales_Rep") or "").strip()
        supplied_id = str(score_row.get("Opportunity_ID_15") or "").strip()
        selected_record: dict[str, Any] | None = None
        resolution_source = "scorecard_id"
        candidates: list[dict[str, Any]] = []

        if supplied_id:
            try:
                canonical_id = canonical_salesforce_id(supplied_id)
            except ValueError as exc:
                add_issue(
                    "invalid_scorecard_id",
                    scorecard_file,
                    [int(score_row["source_row"])],
                    supplied_id,
                    str(exc),
                    score_row,
                    "Correct the scorecard Salesforce ID",
                )
                continue
            selected_record = current_by_canonical.get(canonical_id)
            if not selected_record:
                add_issue(
                    "unmatched_scorecard_id",
                    scorecard_file,
                    [int(score_row["source_row"])],
                    supplied_id,
                    "Scorecard Salesforce ID does not match the current Salesforce export",
                    score_row,
                    "Correct the scorecard ID or source export",
                )
                continue
        else:
            scorecard_source_row = int(score_row["source_row"])
            raw_candidates = [
                record
                for record in current_records
                if normalize_name(record["Product"]) == normalize_name(product)
                and normalize_name(record["Sales_Rep"]) == normalize_name(sales_rep)
            ]
            candidates = sorted(
                [
                    {
                        "opportunity_id": record["source_id"],
                        "source_row": int(record["source_row"]),
                        "hospital": record["Health_System"],
                        "product": record["Product"],
                        "sales_rep": record["Sales_Rep"],
                        "name_similarity": name_similarity(shorthand, record["Health_System"]),
                    }
                    for record in raw_candidates
                ],
                key=lambda candidate: (-candidate["name_similarity"], candidate["hospital"]),
            )
            mapping_match = scorecard_mappings_by_row.get(scorecard_source_row)
            mapping_reason = ""
            if mapping_match:
                mapping_position, mapping = mapping_match
                used_scorecard_mapping_positions.add(mapping_position)
                resolution_id = str(mapping["opportunity_id"]).strip()
                mapping_reason = str(mapping["reason"]).strip()
                selected_record = current_by_source_id.get(resolution_id)
                if not selected_record:
                    try:
                        selected_record = current_by_canonical.get(canonical_salesforce_id(resolution_id))
                    except ValueError:
                        selected_record = None
                resolution_source = "config/resolutions.json"
                if not selected_record:
                    add_issue(
                        "invalid_scorecard_mapping",
                        "config/resolutions.json",
                        [scorecard_source_row],
                        resolution_id,
                        "Confirmed scorecard mapping does not match a current Salesforce opportunity",
                        mapping,
                        "Choose a current Salesforce opportunity ID",
                    )
                    continue
                mapping_status = "human_resolved"
            else:
                identity_matches = [
                    record
                    for record in current_records
                    if normalize_name(record["Health_System"]) == normalize_name(shorthand)
                    and normalize_name(record["Product"]) == normalize_name(product)
                    and normalize_name(record["Sales_Rep"]) == normalize_name(sales_rep)
                ]
                selected_record = identity_matches[0] if len(identity_matches) == 1 else None
                resolution_source = "exact_scorecard_identity_match"
                mapping_status = "automatically_resolved" if selected_record else "blocking"
            scorecard_identity = {
                "hospital": shorthand,
                "product": product,
                "sales_rep": sales_rep,
            }
            if not selected_record:
                add_issue(
                    "missing_scorecard_id",
                    scorecard_file,
                    [scorecard_source_row],
                    None,
                    f"{shorthand} has no unique exact Salesforce identity match",
                    {"scorecard_row": scorecard_identity, "candidates": candidates},
                    "Confirm an opportunity ID in config/resolutions.json",
                )
                continue
            if normalize_name(selected_record["Product"]) != normalize_name(product) or normalize_name(
                selected_record["Sales_Rep"]
            ) != normalize_name(sales_rep):
                add_issue(
                    "invalid_scorecard_mapping",
                    scorecard_file,
                    [scorecard_source_row],
                    selected_record["source_id"],
                    "Confirmed mapping does not match the scorecard product and sales representative",
                    {"scorecard": scorecard_identity, "selected": compact_variant(selected_record)},
                    "Choose a matching Salesforce opportunity",
                )
                continue
            add_audit(
                "missing_scorecard_id",
                mapping_status,
                scorecard_file,
                [scorecard_source_row],
                selected_record["source_id"],
                {"hospital": shorthand, "opportunity_id": None},
                {"opportunity_id": selected_record["source_id"]},
                (
                    f"Applied and validated the confirmed Salesforce mapping: {mapping_reason}"
                    if mapping_status == "human_resolved"
                    else "Mapped the blank scorecard ID from an exact hospital, product, and sales rep match"
                ),
                resolution_source,
                None,
            )

        if not selected_record:
            continue
        if selected_record["Product"] != product or selected_record["Sales_Rep"] != sales_rep:
            add_issue(
                "scorecard_identity_mismatch",
                scorecard_file,
                [int(score_row["source_row"])],
                selected_record["source_id"],
                "Scorecard product or sales representative disagrees with Salesforce",
                {"scorecard": score_row, "salesforce": selected_record},
                "Correct the scorecard or Salesforce identity fields",
            )
            continue
        if not is_active_salesforce(selected_record):
            add_issue(
                "inactive_top8_mapping",
                scorecard_file,
                [int(score_row["source_row"])],
                selected_record["source_id"],
                "Top 8 mapping points to a deal that is not active in Q4-2026",
                {"fiscal_period": selected_record["Fiscal_Period"], "stage": selected_record["Stage"]},
                "Correct the mapping or source deal status",
            )
            continue
        try:
            scores = []
            for field in ("Leadership_Score", "Commercial_Score", "Contractual_Score"):
                raw_score = score_row.get(field)
                number = float(raw_score)
                if not number.is_integer():
                    raise ValueError(f"{field} must be a whole number")
                scores.append(int(number))
            gpa, probability = calculate_top8_probability(*scores)
        except (TypeError, ValueError) as exc:
            add_issue(
                "invalid_top8_scores",
                scorecard_file,
                [int(score_row["source_row"])],
                selected_record["source_id"],
                str(exc),
                score_row,
                "Correct the scorecard scores",
            )
            continue
        canonical_id = selected_record["canonical_id"]
        if canonical_id in top8_by_canonical:
            add_issue(
                "duplicate_top8_mapping",
                scorecard_file,
                [int(score_row["source_row"])],
                selected_record["source_id"],
                "Two scorecard rows map to the same Salesforce opportunity",
                score_row,
                "Correct the scorecard mappings",
            )
            continue
        top8_by_canonical[canonical_id] = {
            "scorecard_source_row": int(score_row["source_row"]),
            "rank": int(score_row["Rank"]),
            "hospital_shorthand": shorthand,
            "leadership": scores[0],
            "commercial": scores[1],
            "contractual": scores[2],
            "composite_gpa": gpa,
            "probability": probability,
            "vp_notes": str(score_row.get("VP_Notes") or ""),
        }
        stored_probability = as_number(score_row.get("Win_Probability"), "Win_Probability")
        if stored_probability != probability:
            add_audit(
                "scorecard_probability_mismatch",
                "automatically_resolved",
                scorecard_file,
                [int(score_row["source_row"])],
                selected_record["source_id"],
                {"workbook_probability": stored_probability, "composite_gpa": gpa},
                {"rule_probability": probability},
                "Used the probability required by the written GPA rule",
                "challenge_rule_2",
                selected_record["annual_acv"] * (probability - stored_probability),
            )

    for position, mapping in enumerate(scorecard_mappings):
        if position in valid_scorecard_mapping_positions and position not in used_scorecard_mapping_positions:
            add_issue(
                "stale_scorecard_mapping",
                "config/resolutions.json",
                [int(mapping["source_row"])],
                str(mapping["opportunity_id"]),
                "Scorecard mapping no longer targets a blank-ID scorecard row",
                mapping,
                "Remove or reconfirm the stale mapping",
            )

    deals: list[dict[str, Any]] = []
    for record in current_records:
        if not is_active_salesforce(record):
            continue
        top8 = top8_by_canonical.get(record["canonical_id"])
        if top8:
            bucket = "top_8"
            probability = top8["probability"]
            scores = {
                "leadership": top8["leadership"],
                "commercial": top8["commercial"],
                "contractual": top8["contractual"],
                "composite_gpa": top8["composite_gpa"],
            }
        else:
            bucket = "salesforce_longtail"
            probability = record["stage_probability"]
            scores = None
        deal = {
            "opportunity_id": record["source_id"],
            "canonical_id": record["canonical_id"],
            "source": "salesforce",
            "source_row": int(record["source_row"]),
            "health_system": record["Health_System"],
            "product": record["Product"],
            "sales_rep": record["Sales_Rep"],
            "region": record.get("Region"),
            "fiscal_period": record["Fiscal_Period"],
            "stage": record["Stage"],
            "stage_probability": record["stage_probability"],
            "close_date": record["Close_Date"],
            "original_amount": record["amount"],
            "billing_frequency": record["Billing_Frequency"],
            "annual_acv": record["annual_acv"],
            "forecast_bucket": bucket,
            "top_8_scores": scores,
            "vp_notes": top8["vp_notes"] if top8 else None,
            "probability": probability,
            "expected_value": record["annual_acv"] * probability,
            "source_conflict": None,
        }
        deals.append(deal)

    for record in carebridge_records:
        if record["Billing_Frequency"] == "Monthly":
            active_impact = (
                (record["annual_acv"] - record["amount"]) * record["stage_probability"]
                if is_active_carebridge(record)
                else 0
            )
            add_audit(
                "monthly_to_annual_conversion",
                "automatically_resolved",
                carebridge_file,
                [int(record["source_row"])],
                record["source_id"],
                {"amount": record["amount"], "billing_frequency": "Monthly"},
                {"annual_acv": record["annual_acv"]},
                "Multiplied the monthly amount by 12",
                "challenge_rule_1",
                active_impact,
            )

    for record in carebridge_records:
        if not is_active_carebridge(record):
            continue
        deal = {
            "opportunity_id": record["source_id"],
            "canonical_id": record["canonical_id"],
            "source": "carebridge",
            "source_row": int(record["source_row"]),
            "health_system": record["Company_Name"],
            "product": record["Product"],
            "sales_rep": record["AE_Owner"],
            "region": None,
            "fiscal_period": record["Fiscal_Period"],
            "stage": record["Deal_Stage"],
            "stage_probability": record["stage_probability"],
            "close_date": record["Expected_Close_Date"],
            "original_amount": record["amount"],
            "billing_frequency": record["Billing_Frequency"],
            "annual_acv": record["annual_acv"],
            "forecast_bucket": "carebridge_longtail",
            "top_8_scores": None,
            "vp_notes": None,
            "probability": record["stage_probability"],
            "expected_value": record["annual_acv"] * record["stage_probability"],
            "source_conflict": None,
        }
        deals.append(deal)

    # Reconcile strong cross-source identity matches, or block when their financial values conflict.
    sf_fingerprints: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for deal in deals:
        if deal["source"] != "salesforce":
            continue
        key = (
            normalize_name(deal["health_system"]),
            normalize_name(deal["product"]),
            normalize_name(deal["sales_rep"]),
            deal["close_date"],
        )
        sf_fingerprints[key].append(deal)
    reconciled: list[dict[str, Any]] = []
    for deal in deals:
        if deal["source"] != "carebridge":
            reconciled.append(deal)
            continue
        key = (
            normalize_name(deal["health_system"]),
            normalize_name(deal["product"]),
            normalize_name(deal["sales_rep"]),
            deal["close_date"],
        )
        salesforce_matches = sf_fingerprints.get(key, [])
        if len(salesforce_matches) != 1:
            reconciled.append(deal)
            continue
        salesforce_match = salesforce_matches[0]
        conflicts: dict[str, dict[str, float]] = {}
        for field in ("annual_acv", "probability"):
            if salesforce_match[field] != deal[field]:
                conflicts[field] = {
                    "salesforce": salesforce_match[field],
                    "carebridge": deal[field],
                }
        if conflicts:
            add_issue(
                "potential_cross_source_duplicate",
                f"{current_file} + {carebridge_file}",
                [salesforce_match["source_row"], deal["source_row"]],
                salesforce_match["opportunity_id"],
                "Strong cross-source identity match has conflicting financial values",
                {
                    "salesforce_id": salesforce_match["opportunity_id"],
                    "hubspot_id": deal["opportunity_id"],
                    "matched_identity": {
                        "health_system": normalize_name(deal["health_system"]),
                        "product": normalize_name(deal["product"]),
                        "sales_rep": normalize_name(deal["sales_rep"]),
                        "close_date": deal["close_date"],
                    },
                    "conflicts": conflicts,
                },
                "Confirm whether the records are the same deal and select the correct financial values",
            )
            reconciled.append(deal)
            continue
        add_audit(
            "cross_source_duplicate",
            "automatically_resolved",
            f"{current_file} + {carebridge_file}",
            [salesforce_match["source_row"], deal["source_row"]],
            salesforce_match["opportunity_id"],
            {
                "salesforce_id": salesforce_match["opportunity_id"],
                "hubspot_id": deal["opportunity_id"],
                "annual_acv": deal["annual_acv"],
            },
            {"retained_source": "salesforce", "calculation_record_count": 1},
            "Counted the equivalent cross-system deal once using the Salesforce record",
            "pipeline_identity_rule",
            -deal["expected_value"],
        )
    deals = reconciled

    # Surface stale close dates as warnings without changing eligibility.
    for deal in deals:
        try:
            close_date = datetime.strptime(deal["close_date"], "%Y-%m-%d").date()
        except (TypeError, ValueError):
            add_audit(
                "invalid_close_date",
                "warning",
                current_file if deal["source"] == "salesforce" else carebridge_file,
                [deal["source_row"]],
                deal["opportunity_id"],
                {"close_date": deal["close_date"]},
                None,
                "Kept the deal because fiscal period and stage determine eligibility",
                "eligibility_rule",
                0,
            )
            continue
        if close_date < REPORTING_DATE:
            add_audit(
                "stale_close_date",
                "warning",
                current_file if deal["source"] == "salesforce" else carebridge_file,
                [deal["source_row"]],
                deal["opportunity_id"],
                {"close_date": deal["close_date"], "reporting_date": REPORTING_DATE.isoformat()},
                {"deal_retained": True},
                "Kept the deal because fiscal period and stage determine eligibility",
                "eligibility_rule",
                0,
            )

    review = build_review()
    if issues:
        return review, None

    week_over_week = build_week_over_week(current_records, prior_records)
    deals.sort(key=lambda deal: (deal["forecast_bucket"] != "top_8", deal["source"], deal["health_system"]))
    top8_value = sum(deal["expected_value"] for deal in deals if deal["forecast_bucket"] == "top_8")
    sf_longtail_value = sum(
        deal["expected_value"] for deal in deals if deal["forecast_bucket"] == "salesforce_longtail"
    )
    carebridge_value = sum(
        deal["expected_value"] for deal in deals if deal["forecast_bucket"] == "carebridge_longtail"
    )
    total = top8_value + sf_longtail_value + carebridge_value
    forecast = {
        "schema_version": 1,
        "reporting_date": REPORTING_DATE.isoformat(),
        "generated_at": generated_at,
        "board_target": BOARD_TARGET,
        "summary": {
            "total_expected_sales": total,
            "top_8_expected_value": top8_value,
            "salesforce_longtail_expected_value": sf_longtail_value,
            "carebridge_longtail_expected_value": carebridge_value,
            "variance_to_target": total - BOARD_TARGET,
            "baseline_status": "valid",
        },
        "deals": deals,
        "week_over_week": week_over_week,
        "audit_log": audits,
    }
    return review, forecast


def atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def run_pipeline(root: Path) -> int:
    review, forecast = build_outputs(root)
    atomic_write_json(root / OUTPUT_DIR / "review.json", review)
    if forecast is None:
        print(
            f"Pipeline blocked by {review['blocking_issue_count']} issue(s). "
            "See data/output/review.json.",
            file=sys.stderr,
        )
        return 1
    atomic_write_json(root / OUTPUT_DIR / "forecast.json", forecast)
    print(
        "Forecast generated: "
        f"${forecast['summary']['total_expected_sales']:,.0f} total expected sales, "
        f"{len(forecast['deals'])} active deals, "
        f"{len(forecast['audit_log'])} audit entries."
    )
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="Project root containing data/input, config/input_files.json, and config/resolutions.json",
    )
    args = parser.parse_args()
    try:
        return run_pipeline(args.root.resolve())
    except Exception as exc:  # Technical failures must not replace generated outputs.
        print(f"Pipeline failed: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
