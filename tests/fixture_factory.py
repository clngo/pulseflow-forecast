from __future__ import annotations

import csv
import json
import shutil
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


FIXTURES_ROOT = Path(__file__).resolve().parent / "fixtures"
BASELINE_ROOT = FIXTURES_ROOT / "baseline"
SCENARIOS_PATH = FIXTURES_ROOT / "scenarios.json"

DATASET_PATHS = {
    "current": Path("data/input/1_sfdc_pipeline_current.csv"),
    "prior": Path("data/input/2_sfdc_pipeline_prior_week.csv"),
    "scorecard": Path("data/input/3_vp_deal_scorecard.xlsx"),
    "carebridge": Path("data/input/4_carebridge_hubspot_export.csv"),
}


def load_scenarios() -> dict[str, dict[str, Any]]:
    return json.loads(SCENARIOS_PATH.read_text(encoding="utf-8"))


def scenario_definition(scenario_id: str) -> dict[str, Any]:
    scenarios = load_scenarios()
    try:
        return scenarios[scenario_id]
    except KeyError as exc:
        raise KeyError(f"Unknown fixture scenario: {scenario_id}") from exc


def _matches(row: dict[str, Any], expected: dict[str, Any]) -> bool:
    return all(row.get(field) == value for field, value in expected.items())


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows({header: row.get(header, "") for header in headers} for row in rows)


def _matching_row(rows: list[dict[str, Any]], match: dict[str, Any], dataset: str) -> dict[str, Any]:
    matches = [row for row in rows if _matches(row, match)]
    if len(matches) != 1:
        raise ValueError(f"Expected one {dataset} fixture row matching {match}, found {len(matches)}")
    return matches[0]


def _apply_csv_operation(path: Path, dataset: str, operation: dict[str, Any]) -> None:
    headers, rows = _read_csv(path)
    action = operation["action"]

    if action == "set":
        row = _matching_row(rows, operation["match"], dataset)
        row.update(operation["values"])
    elif action == "copy_row":
        matching_rows = [row for row in rows if _matches(row, operation["match"])]
        if not matching_rows:
            raise ValueError(f"No {dataset} fixture row matches {operation['match']}")
        source = matching_rows[0]
        copied = deepcopy(source)
        copied.update(operation.get("values", {}))
        rows.append(copied)
    elif action == "remove_column":
        column = operation["column"]
        if column not in headers:
            raise ValueError(f"Fixture column is already absent: {column}")
        headers.remove(column)
        for row in rows:
            row.pop(column, None)
    else:
        raise ValueError(f"Unsupported CSV fixture action: {action}")

    _write_csv(path, headers, rows)


def _apply_scorecard_operation(path: Path, operation: dict[str, Any]) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        headers = {
            sheet.cell(row=4, column=column).value: column
            for column in range(1, sheet.max_column + 1)
        }
        matches: list[int] = []
        for row_number in range(5, sheet.max_row + 1):
            row = {
                header: sheet.cell(row=row_number, column=column).value
                for header, column in headers.items()
            }
            if _matches(row, operation["match"]):
                matches.append(row_number)
        if len(matches) != 1:
            raise ValueError(
                f"Expected one scorecard fixture row matching {operation['match']}, found {len(matches)}"
            )
        action = operation["action"]
        if action == "set":
            target_row = matches[0]
        elif action == "copy_row":
            target_row = sheet.max_row + 1
            for column in range(1, sheet.max_column + 1):
                sheet.cell(row=target_row, column=column).value = sheet.cell(
                    row=matches[0], column=column
                ).value
        else:
            raise ValueError(f"Unsupported scorecard fixture action: {action}")
        for field, value in operation.get("values", {}).items():
            sheet.cell(row=target_row, column=headers[field]).value = value
        workbook.save(path)
    finally:
        workbook.close()


def make_scenario_dataset(root: Path, scenario_id: str) -> Path:
    """Create an isolated working copy from committed, inspectable fixture definitions."""
    scenario = scenario_definition(scenario_id)
    shutil.copytree(BASELINE_ROOT, root)

    for operation in scenario["operations"]:
        dataset = operation["dataset"]
        path = root / DATASET_PATHS[dataset]
        if dataset == "scorecard":
            _apply_scorecard_operation(path, operation)
        else:
            _apply_csv_operation(path, dataset, operation)
    return root


def find_audit(forecast: dict[str, Any], issue_type: str, record_id: str) -> dict[str, Any]:
    return next(
        entry
        for entry in forecast["audit_log"]
        if entry["type"] == issue_type and entry["record_id"] == record_id
    )


def find_issue(review: dict[str, Any], issue_type: str, record_id: str) -> dict[str, Any]:
    return next(
        issue
        for issue in review["issues"]
        if issue["type"] == issue_type and issue["record_id"] == record_id
    )
